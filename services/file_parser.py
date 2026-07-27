"""
services/file_parser.py — PDF & DOCX Text Extraction
======================================================
Extracts raw text from uploaded resume files.
Supports: .pdf (via pdfplumber), .docx (via python-docx)
"""

import os
import pdfplumber
import docx


def extract_text_from_pdf(path: str) -> str:
    """Extract all text from a PDF file using pdfplumber."""
    text_parts = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        raise ValueError(f"Failed to read PDF '{os.path.basename(path)}': {e}")
    return "\n".join(text_parts)


def extract_text_from_docx(path: str) -> str:
    """Extract all text from a DOCX file using python-docx."""
    try:
        doc   = docx.Document(path)
        lines = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract table cell text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        lines.append(cell.text.strip())
        return "\n".join(lines)
    except Exception as e:
        raise ValueError(f"Failed to read DOCX '{os.path.basename(path)}': {e}")


def parse_resume(path: str) -> str:
    """
    Dispatcher: choose parser based on file extension.
    Returns the cleaned raw text of the resume.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        raw = extract_text_from_pdf(path)
    elif ext == ".docx":
        raw = extract_text_from_docx(path)
    else:
        raise ValueError(f"Unsupported file type: '{ext}'. Use PDF or DOCX.")

    if not raw or not raw.strip():
        raise ValueError("Could not extract any text from the resume. "
                         "Ensure it is not a scanned image-only PDF.")
    return raw.strip()
